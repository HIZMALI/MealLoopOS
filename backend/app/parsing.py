"""Conservative, bounded import. Missing measurements remain null; rows never sort."""
import csv
import io
import math
import re
import zipfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 64

ALIASES = {
    "date": "date", "menu": "menu", "menu identifier": "menu", "diners": "diners",
    "diner count": "diners", "pw kg": "plate_waste_kg", "pw": "plate_waste_kg",
    "plate_waste_kg": "plate_waste_kg", "plate waste kg": "plate_waste_kg",
    "ksw kg": "kitchen_serving_waste_kg", "ksw": "kitchen_serving_waste_kg",
    "kitchen_serving_waste_kg": "kitchen_serving_waste_kg",
    "kitchen serving waste kg": "kitchen_serving_waste_kg",
    "pw g": "reported_plate_waste_g", "plate_waste_g": "reported_plate_waste_g",
    "pw/g/diner": "reported_plate_g_per_diner", "reported_plate_g_per_diner": "reported_plate_g_per_diner",
    "ksw/g/diner": "reported_kitchen_g_per_diner", "reported_kitchen_g_per_diner": "reported_kitchen_g_per_diner",
}
NUMERIC = ("menu", "diners", "plate_waste_kg", "kitchen_serving_waste_kg", "reported_plate_waste_g",
           "reported_plate_g_per_diner", "reported_kitchen_g_per_diner")


class ImportProblem(ValueError):
    def __init__(self, message: str, errors=None):
        super().__init__(message)
        self.errors = errors or []


def blank(value):
    return value is None or (isinstance(value, str) and value.strip().lower() in ("", "na", "n/a", "null"))


def read_table(content: bytes, filename: str):
    suffix = Path(filename).suffix.lower()
    if suffix not in (".xlsx", ".csv"):
        raise ImportProblem("Only .xlsx and UTF-8 .csv files are supported.")
    if not content or len(content) > MAX_BYTES:
        raise ImportProblem("Upload a non-empty file no larger than 10 MiB.")
    try:
        if suffix == ".csv":
            text = content.decode("utf-8-sig")
            delimiter = ";" if text.splitlines()[0].count(";") > text.splitlines()[0].count(",") else ","
            table = list(csv.reader(io.StringIO(text), delimiter=delimiter, strict=True))
            tables = [("CSV", table)]
        else:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                info = archive.infolist()
                if len(info) > 1000 or sum(x.file_size for x in info) > 50 * 1024 * 1024:
                    raise ImportProblem("Workbook expands beyond the safe 50 MiB / 1,000-entry limit.")
                if any("vbaProject" in x.filename for x in info):
                    raise ImportProblem("Macro-enabled workbooks are not supported.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                if len(workbook.worksheets) > 20:
                    raise ImportProblem("Workbook exceeds the 20-sheet limit.")
                tables = []
                for sheet in workbook:
                    if sheet.max_row > 20_000 or sheet.max_column > MAX_COLUMNS:
                        raise ImportProblem("Workbook exceeds the 20,000 physical rows / 64 columns limit.")
                    tables.append((sheet.title, list(sheet.values)))
            finally:
                workbook.close()
        candidates = []
        for name, table in tables:
            nonempty = [(i + 1, list(row)) for i, row in enumerate(table) if any(not blank(x) for x in row)]
            if not nonempty:
                continue
            header = nonempty[0][1]
            mapped = [ALIASES.get(re.sub(r"\s+", " ", str(x).strip().lower())) if not blank(x) else None for x in header]
            if "diners" in mapped and ("plate_waste_kg" in mapped or "reported_plate_waste_g" in mapped):
                candidates.append((name, nonempty, mapped))
        if len(candidates) != 1:
            raise ImportProblem("Expected exactly one data sheet with a first non-empty header row containing DINERS and PW kg (or PW g). See the CSV template. Multiple matching sheets must be uploaded separately.")
        name, rows, mapped = candidates[0]
        if "kitchen_serving_waste_kg" not in mapped:
            raise ImportProblem("Missing required kitchen/serving mass column: KSW kg or kitchen_serving_waste_kg. Empty cells are allowed; a missing column is not.")
        recognized = [x for x in mapped if x]
        if len(recognized) != len(set(recognized)):
            raise ImportProblem("Two columns map to the same field. Keep one column per measurement and unit.")
        if len(rows) < 2 or len(rows) - 1 > MAX_ROWS or any(len(row) > MAX_COLUMNS for _, row in rows):
            raise ImportProblem("Expected 1–10,000 data rows and no more than 64 columns.")
        return name, rows, mapped, [n for n, _ in tables]
    except ImportProblem:
        raise
    except Exception as exc:
        raise ImportProblem("Could not read this file. Supply a valid, unencrypted .xlsx or UTF-8 CSV.") from exc


def number(value, field, source_row):
    if blank(value):
        return None
    try:
        if isinstance(value, bool):
            raise ValueError()
        # Decimal commas are accepted, but thousands separators are deliberately unsupported.
        numeric = float(str(value).strip().replace(",", "."))
        if not math.isfinite(numeric) or numeric < 0:
            raise ValueError()
        if field in ("diners", "menu"):
            if not numeric.is_integer() or numeric > 1_000_000:
                raise ValueError()
            if field == "menu" and not 1 <= numeric <= 25:
                raise ValueError()
            return int(numeric)
        if numeric > 1_000_000:
            raise ValueError()
        return numeric
    except (ValueError, TypeError, OverflowError):
        raise ImportProblem("Invalid numeric data.", [{"source_row": source_row, "field": field,
            "message": "Expected a finite nonnegative number (integer for diners; menu 1–25), at most 1,000,000."}])


def parse_date(value, source_row):
    if blank(value):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ImportProblem("Invalid date. Use ISO YYYY-MM-DD or DD/MM/YYYY; dates are optional.",
                        [{"source_row": source_row, "field": "date", "message": "Unrecognized date; numeric serials in CSV are not guessed."}])


def parse_records(content: bytes, filename: str, baseline_days: int | None = None):
    sheet, table, mapping, sheets = read_table(content, filename)
    if baseline_days is not None and not 0 <= baseline_days <= len(table) - 1:
        raise ImportProblem("Baseline day count must be between zero and the number of imported records.")
    records, errors, seen, seen_dates = [], [], set(), set()
    for day_index, (source_row, values) in enumerate(table[1:], 1):
        raw = {field: values[j] if j < len(values) else None for j, field in enumerate(mapping) if field}
        try:
            record = {field: number(raw.get(field), field, source_row) for field in NUMERIC}
            record["date"] = parse_date(raw.get("date"), source_row)
            issues = []
            if "plate_waste_kg" not in mapping and record["reported_plate_waste_g"] is not None:
                record["plate_waste_kg"] = record["reported_plate_waste_g"] / 1000
                issues.append("plate_mass_converted_from_grams")
            if record["plate_waste_kg"] is None and record["kitchen_serving_waste_kg"] is None:
                raise ImportProblem("A record must contain at least one waste mass measurement.",
                    [{"source_row": source_row, "field": "waste_mass", "message": "Both waste masses are missing; row retained only after correcting the input."}])
            for field in ("plate_waste_kg", "kitchen_serving_waste_kg", "diners", "menu", "date"):
                if record[field] is None:
                    issues.append(f"missing_{field}")
            p, g, k, n = (record[x] for x in ("plate_waste_kg", "reported_plate_waste_g", "kitchen_serving_waste_kg", "diners"))
            if p is not None and g is not None and not math.isclose(p * 1000, g, abs_tol=0.5):
                issues.append("plate_kg_g_conflict")
            if n == 0:
                issues.append("zero_diners_per_diner_unavailable")
            if n and n > 0:
                for total, per, code in ((p, record["reported_plate_g_per_diner"], "plate_per_diner_conflict"),
                                         (k, record["reported_kitchen_g_per_diner"], "kitchen_per_diner_conflict")):
                    if total is not None and per is not None and not math.isclose(total * 1000 / n, per, abs_tol=0.05):
                        issues.append(code)
            signature = tuple(record.items())
            if signature in seen:
                issues.append("duplicate_measurements_preserved")
            seen.add(signature)
            if record["date"] is not None:
                if record["date"] in seen_dates:
                    issues.append("duplicate_date_preserved")
                seen_dates.add(record["date"])
            record.update(day_index=day_index, source_row=source_row, issues=issues,
                phase="unassigned" if baseline_days is None else "baseline" if day_index <= baseline_days else "intervention")
            records.append(record)
        except ImportProblem as exc:
            errors.extend(exc.errors)
    if errors:
        raise ImportProblem(f"Import rejected: {len(errors)} invalid row(s). No records were saved. Correct the input and retry.", errors[:50])
    dates = [r["date"] for r in records if r["date"]]
    date_ordered = all(a <= b for a, b in zip(dates, dates[1:]))
    counts = dict(Counter(issue for r in records for issue in r["issues"]))
    inspection = {"sheet": sheet, "sheets": sheets, "headers": [str(x) for x in table[0][1] if not blank(x)],
        "header_mapping": {str(h): field for h, field in zip(table[0][1], mapping) if field},
        "ignored_columns": [str(h) for h, field in zip(table[0][1], mapping) if not field and not blank(h)],
        "issue_counts": counts, "flagged_rows": sum(bool(r["issues"]) for r in records),
        "date_count": len(dates), "date_ordered": date_ordered, "order": "source row order",
        "missing_counts": {field: sum(r[field] is None for r in records) for field in ("date", *NUMERIC)},
        "date_start": str(min(dates)) if dates else None, "date_end": str(max(dates)) if dates else None}
    return records, inspection

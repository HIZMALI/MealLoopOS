"""Inspect the actual workbook, including cells, formats, missingness and conflicts."""
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from app.database import ROOT
from app.parsing import parse_records

path = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/raw/Food waste data school Finland 2024-2025.xlsx"
content = path.read_bytes()
rows, inspection = parse_records(content, path.name, 135)
book = load_workbook(path, data_only=True, read_only=True)
sheet = book.active
inspection["file_sha256"] = hashlib.sha256(content).hexdigest()
inspection["file_md5"] = hashlib.md5(content).hexdigest()
inspection["physical_dimensions"] = {"rows": sheet.max_row, "columns": sheet.max_column}
inspection["column_formats"] = {str(sheet.cell(1, j).value): dict(Counter(sheet.cell(i, j).number_format for i in range(2, 157))) for j in range(1, 9)}
inspection["flagged_records"] = [r for r in rows if r["issues"]]
inspection["record_count"] = len(rows)
book.close()
print(json.dumps(inspection, indent=2, default=str, ensure_ascii=True))
